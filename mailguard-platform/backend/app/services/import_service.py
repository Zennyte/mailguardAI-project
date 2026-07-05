import csv
import io
import json

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import (
    EmailMessage, CmsPage, UserFeedback, Setting, Notification,
    ImportJob, ScanResult, ScanRequest,
)
from app.schemas.data_transfer import ImportSummary

IMPORT_ENTITIES = ["email_messages", "cms_pages", "user_feedback", "settings", "notifications"]

VALID_LABELS = ["safe", "spam", "phishing"]


def import_entity(db: Session, user_id: int, entity: str,
                  filename: str, content: bytes) -> ImportSummary:
    if entity not in IMPORT_ENTITIES:
        raise HTTPException(status_code=400,
                            detail=f"Unknown entity. Allowed: {', '.join(IMPORT_ENTITIES)}")

    rows = _read_rows(filename, content)

    if entity == "email_messages":
        imported, skipped = _import_email_messages(db, user_id, rows)
    elif entity == "cms_pages":
        imported, skipped = _import_cms_pages(db, user_id, rows)
    elif entity == "user_feedback":
        imported, skipped = _import_user_feedback(db, user_id, rows)
    elif entity == "settings":
        imported, skipped = _import_settings(db, user_id, rows)
    else:
        imported, skipped = _import_notifications(db, user_id, rows)

    # Regjistrojme importin ne tabelen import_jobs
    db.add(ImportJob(user_id=user_id, status="completed",
                     total_rows=len(rows), processed_rows=imported))
    db.commit()

    return ImportSummary(
        entity=entity,
        imported_count=imported,
        skipped_count=skipped,
        message=f"Imported {imported} rows, skipped {skipped} invalid rows.",
    )


def _read_rows(filename: str, content: bytes) -> list:
    # Formati njihet nga prapashtesa e skedarit
    name = (filename or "").lower()
    try:
        if name.endswith(".csv"):
            text = content.decode("utf-8-sig")
            return [dict(row) for row in csv.DictReader(io.StringIO(text))]
        if name.endswith(".json"):
            data = json.loads(content)
            if not isinstance(data, list):
                raise ValueError("JSON must be a list of objects")
            return data
        if name.endswith(".xlsx"):
            sheet = load_workbook(io.BytesIO(content), read_only=True).active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(h) for h in next(iterator, [])]
            return [dict(zip(headers, row)) for row in iterator]
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the uploaded file.")

    raise HTTPException(status_code=400, detail="Unsupported file type. Use .csv, .json or .xlsx")


def _text(row: dict, key: str) -> str:
    value = row.get(key)
    return str(value).strip() if value is not None else ""


def _to_bool(value) -> bool:
    return str(value).strip().lower() in ("true", "1", "yes")


def _import_email_messages(db, user_id, rows):
    imported = skipped = 0
    for row in rows:
        body = _text(row, "body")
        if not body:
            skipped += 1
            continue
        db.add(EmailMessage(
            user_id=user_id,
            subject=_text(row, "subject")[:500],
            body_preview=body[:500],
            source_type="imported",
        ))
        imported += 1
    return imported, skipped


def _import_cms_pages(db, user_id, rows):
    imported = skipped = 0
    for row in rows:
        title = _text(row, "title")
        slug = _text(row, "slug")
        exists = slug and db.query(CmsPage).filter(CmsPage.slug == slug).first()
        if not title or not slug or exists:
            skipped += 1
            continue
        db.add(CmsPage(
            title=title[:200],
            slug=slug[:200],
            is_published=_to_bool(row.get("is_published")),
            created_by=user_id,
        ))
        db.flush()  # qe slug-u i ri te njihet nga rreshtat pasues
        imported += 1
    return imported, skipped


def _import_user_feedback(db, user_id, rows):
    imported = skipped = 0
    for row in rows:
        try:
            scan_result_id = int(row.get("scan_result_id"))
        except (TypeError, ValueError):
            skipped += 1
            continue

        # Feedback lejohet vetem per skanimet e vetat
        scan_result = (
            db.query(ScanResult)
            .join(ScanRequest, ScanRequest.id == ScanResult.scan_request_id)
            .filter(ScanResult.id == scan_result_id, ScanRequest.user_id == user_id)
            .first()
        )
        correct_label = _text(row, "correct_label")
        if scan_result is None or (correct_label and correct_label not in VALID_LABELS):
            skipped += 1
            continue

        db.add(UserFeedback(
            scan_result_id=scan_result_id,
            user_id=user_id,
            correct_label=correct_label or None,
            comment=_text(row, "comment") or None,
        ))
        imported += 1
    return imported, skipped


def _import_settings(db, user_id, rows):
    imported = skipped = 0
    for row in rows:
        key = _text(row, "key") or _text(row, "setting_key")
        exists = key and db.query(Setting).filter(Setting.setting_key == key).first()
        if not key or exists:
            # Konfigurimet ekzistuese nuk mbishkruhen
            skipped += 1
            continue
        db.add(Setting(
            setting_key=key[:100],
            setting_value=_text(row, "value") or _text(row, "setting_value"),
            description=_text(row, "description")[:255] or None,
            created_by=user_id,
        ))
        db.flush()
        imported += 1
    return imported, skipped


def _import_notifications(db, user_id, rows):
    imported = skipped = 0
    for row in rows:
        title = _text(row, "title")
        if not title:
            skipped += 1
            continue
        db.add(Notification(
            user_id=user_id,
            title=title[:200],
            message=_text(row, "message") or None,
            notification_type=_text(row, "type") or _text(row, "notification_type") or None,
        ))
        imported += 1
    return imported, skipped
